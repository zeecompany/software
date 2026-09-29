"""Employee master list helpers for the whole AURCO system."""
from __future__ import annotations

import csv
import datetime as _dt
import io
import re
from pathlib import Path
from typing import Any, Sequence

from .database import Database

COLUMNS = [
    ("employee_id", "Employee ID"),
    ("name", "Name"),
    ("designation", "Designation"),
    ("iqama_id", "Iqama ID"),
    ("date_of_joining", "Date of Joining"),
    ("nationality", "Nationality"),
    ("contract_workhours", "Contract WorkHours"),
    ("division", "Division/Department"),
    ("current_project", "Current Project"),
    ("location", "Location"),
]
LABELS = dict(COLUMNS)
ALL_FIELDS = COLUMNS

HEADER_MAP = {
    "employeeid": "employee_id",
    "employeecode": "employee_id",
    "empid": "employee_id",
    "empcode": "employee_id",
    "code": "employee_id",
    "name": "name",
    "employeename": "name",
    "staffname": "name",
    "designation": "designation",
    "position": "designation",
    "jobtitle": "designation",
    "iqamaid": "iqama_id",
    "iqama": "iqama_id",
    "idnumber": "iqama_id",
    "dateofjoining": "date_of_joining",
    "joiningdate": "date_of_joining",
    "doj": "date_of_joining",
    "nationality": "nationality",
    "country": "nationality",
    "contractworkhours": "contract_workhours",
    "workhours": "contract_workhours",
    "hours": "contract_workhours",
    "divisiondepartment": "division",
    "division": "division",
    "department": "division",
    "dept": "division",
    "currentproject": "current_project",
    "project": "current_project",
    "site": "current_project",
    "location": "location",
    "camp": "location",
}

_DATE_PATTERNS = (
    "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d.%m.%Y",
    "%Y/%m/%d", "%d-%b-%Y", "%d %b %Y", "%b %d, %Y", "%d-%b-%y",
    "%d-%m-%y", "%d/%m/%y", "%Y-%m-%d %H:%M:%S",
)


def _now() -> str:
    return _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def norm_key(text: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(text or "").lower())


def to_date(v: Any) -> str:
    if v in (None, ""):
        return ""
    if isinstance(v, _dt.datetime):
        return v.date().isoformat()
    if isinstance(v, _dt.date):
        return v.isoformat()
    t = str(v).strip()
    if not t or t in ("-", "N/A", "n/a"):
        return ""
    for f in _DATE_PATTERNS:
        try:
            d = _dt.datetime.strptime(t, f).date()
            if d.year < 100:
                d = d.replace(year=d.year + 2000)
            return d.isoformat()
        except ValueError:
            continue
    try:
        n = float(t)
        if 20000 < n < 60000:
            return (_dt.date(1899, 12, 30) + _dt.timedelta(days=int(n))).isoformat()
    except ValueError:
        pass
    return t


def save_employee(db: Database, data: dict[str, Any], employee_row_id: int | None = None) -> int:
    emp_id = str(data.get("employee_id") or "").strip()
    name = str(data.get("name") or "").strip()
    if not emp_id:
        raise ValueError("Employee ID is required.")
    if not name:
        raise ValueError("Employee name is required.")
    vals = {
        "employee_id": emp_id,
        "name": name,
        "designation": str(data.get("designation") or "").strip(),
        "iqama_id": str(data.get("iqama_id") or "").strip(),
        "date_of_joining": to_date(data.get("date_of_joining") or ""),
        "nationality": str(data.get("nationality") or "").strip(),
        "contract_workhours": str(data.get("contract_workhours") or "").strip(),
        "division": str(data.get("division") or "").strip(),
        "current_project": str(data.get("current_project") or "").strip(),
        "location": str(data.get("location") or "").strip(),
    }
    params = [vals["employee_id"]]
    sql = "SELECT id FROM employees WHERE employee_id=?"
    if employee_row_id is not None:
        sql += " AND id<>?"
        params.append(int(employee_row_id))
    if db.one(sql, tuple(params)):
        raise ValueError(f"Employee ID '{vals['employee_id']}' already exists.")
    if vals["iqama_id"]:
        params = [vals["iqama_id"]]
        sql = "SELECT id FROM employees WHERE iqama_id=?"
        if employee_row_id is not None:
            sql += " AND id<>?"
            params.append(int(employee_row_id))
        if db.one(sql, tuple(params)):
            raise ValueError(f"Iqama ID '{vals['iqama_id']}' already exists.")
    try:
        if employee_row_id is None:
            cur = db.execute(
                """INSERT INTO employees(employee_id,name,designation,iqama_id,date_of_joining,
                     nationality,contract_workhours,division,current_project,location,updated_at)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                (vals["employee_id"], vals["name"], vals["designation"], vals["iqama_id"],
                 vals["date_of_joining"], vals["nationality"], vals["contract_workhours"],
                 vals["division"], vals["current_project"], vals["location"], _now()),
            )
            new_id = int(cur.lastrowid)
            db.commit()
            db.audit("CREATED", "employee", str(new_id), vals["employee_id"])
            return new_id
        db.execute(
            """UPDATE employees SET employee_id=?, name=?, designation=?, iqama_id=?,
                 date_of_joining=?, nationality=?, contract_workhours=?, division=?,
                 current_project=?, location=?, updated_at=?
               WHERE id=?""",
            (vals["employee_id"], vals["name"], vals["designation"], vals["iqama_id"],
             vals["date_of_joining"], vals["nationality"], vals["contract_workhours"],
             vals["division"], vals["current_project"], vals["location"], _now(), int(employee_row_id)),
        )
        db.commit()
        db.audit("EDITED", "employee", str(employee_row_id), vals["employee_id"])
        return int(employee_row_id)
    except Exception:
        db.rollback()
        raise


def delete_employee(db: Database, employee_row_id: int) -> None:
    row = db.one("SELECT employee_id FROM employees WHERE id=?", (employee_row_id,))
    if row is None:
        return
    db.execute("DELETE FROM employees WHERE id=?", (employee_row_id,))
    db.commit()
    db.audit("DELETED", "employee", str(employee_row_id), str(row[0]))


def list_employees(db: Database, text: str = "", division: str = "", project: str = "",
                   location: str = "") -> list[dict[str, Any]]:
    sql = "SELECT * FROM employees WHERE 1=1"
    p: list[Any] = []
    if division:
        sql += " AND division=?"
        p.append(division)
    if project:
        sql += " AND current_project=?"
        p.append(project)
    if location:
        sql += " AND location=?"
        p.append(location)
    if text.strip():
        like = f"%{text.strip()}%"
        sql += (" AND (employee_id LIKE ? OR name LIKE ? OR designation LIKE ? OR iqama_id LIKE ?"
                " OR division LIKE ? OR current_project LIKE ? OR location LIKE ? OR nationality LIKE ?)")
        p += [like] * 8
    sql += " ORDER BY name COLLATE NOCASE, employee_id COLLATE NOCASE"
    return [dict(r) for r in db.query(sql, p)]


def distinct_values(db: Database, field: str) -> list[str]:
    if field not in {"division", "current_project", "location", "designation", "nationality"}:
        return []
    rows = db.query(f"SELECT DISTINCT {field} FROM employees WHERE COALESCE({field},'')<>'' ORDER BY {field}")
    return [str(r[0]) for r in rows if str(r[0] or "").strip()]


def find_employee(db: Database, employee_id: str = "", iqama_id: str = "", name: str = "") -> dict[str, Any] | None:
    emp_id = str(employee_id or "").strip()
    iqama = str(iqama_id or "").strip()
    nm = str(name or "").strip()
    row = None
    if emp_id:
        row = db.one("SELECT * FROM employees WHERE employee_id=?", (emp_id,))
    if row is None and iqama:
        row = db.one("SELECT * FROM employees WHERE iqama_id=?", (iqama,))
    if row is None and nm:
        row = db.one("SELECT * FROM employees WHERE name=?", (nm,))
    return dict(row) if row else None


def employee_names(db: Database) -> list[str]:
    return [str(r[0]) for r in db.query("SELECT name FROM employees WHERE COALESCE(name,'')<>'' ORDER BY name")]


def employee_ids(db: Database) -> list[str]:
    return [str(r[0]) for r in db.query("SELECT employee_id FROM employees WHERE COALESCE(employee_id,'')<>'' ORDER BY employee_id")]


def iqama_ids(db: Database) -> list[str]:
    return [str(r[0]) for r in db.query("SELECT iqama_id FROM employees WHERE COALESCE(iqama_id,'')<>'' ORDER BY iqama_id")]


def sniff(text: str) -> tuple[list[str], list[list[str]]]:
    raw = [ln for ln in text.replace("\r\n", "\n").replace("\r", "\n").split("\n") if ln.strip()]
    if not raw:
        return [], []
    if "\t" in raw[0]:
        rows = [ln.split("\t") for ln in raw]
    elif raw[0].count(",") >= 2:
        rows = list(csv.reader(io.StringIO("\n".join(raw))))
    elif raw[0].count("|") >= 2:
        rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in raw]
    else:
        rows = [re.split(r"\s{2,}", ln.strip()) for ln in raw]
    rows = [[str(c).strip() for c in r] for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return [], []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    head = rows[0]
    if sum(1 for c in head if norm_key(c) in HEADER_MAP) >= max(2, min(4, width)):
        return head, rows[1:]
    return [f"Column {i + 1}" for i in range(width)], rows


def read_file(path: str | Path) -> tuple[list[str], list[list[Any]]]:
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(p, data_only=True, read_only=True)
        ws = wb.active
        data = [[("" if c is None else c) for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        data = [r for r in data if any(str(c).strip() for c in r)]
        if not data:
            return [], []
        head = [str(c).strip() for c in data[0]]
        if sum(1 for c in head if norm_key(c) in HEADER_MAP) >= 2:
            return head, data[1:]
        return [f"Column {i + 1}" for i in range(len(head))], data
    return sniff(p.read_text(encoding="utf-8", errors="ignore"))


def auto_map(headers: Sequence[str]) -> dict[int, str]:
    out: dict[int, str] = {}
    used: set[str] = set()
    for i, h in enumerate(headers):
        f = HEADER_MAP.get(norm_key(h))
        if f and f not in used:
            out[i] = f
            used.add(f)
    return out


def preview(headers: Sequence[str], rows: Sequence[Sequence[Any]],
            mapping: dict[int, str], defaults: dict | None = None) -> list[dict[str, Any]]:
    defaults = defaults or {}
    out: list[dict[str, Any]] = []
    for row in rows:
        rec = {f: "" for f, _ in ALL_FIELDS}
        rec.update({k: v for k, v in defaults.items() if v not in (None, "")})
        for i, field in mapping.items():
            if i < len(row):
                rec[field] = row[i]
        rec["employee_id"] = str(rec.get("employee_id") or "").strip()
        rec["name"] = str(rec.get("name") or "").strip()
        rec["designation"] = str(rec.get("designation") or "").strip()
        rec["iqama_id"] = str(rec.get("iqama_id") or "").strip()
        rec["date_of_joining"] = to_date(rec.get("date_of_joining") or "")
        rec["nationality"] = str(rec.get("nationality") or "").strip()
        rec["contract_workhours"] = str(rec.get("contract_workhours") or "").strip()
        rec["division"] = str(rec.get("division") or "").strip()
        rec["current_project"] = str(rec.get("current_project") or "").strip()
        rec["location"] = str(rec.get("location") or "").strip()
        if not rec["employee_id"] or not rec["name"]:
            continue
        out.append(rec)
    return out


def template_rows() -> tuple[list[str], list[list[Any]]]:
    cols = [lbl for _, lbl in COLUMNS]
    return cols, [
        ["EMP-1001", "Ahmed Salem", "Surveyor", "2456677889", "2024-01-15", "Indian", "10", "Survey", "Hajar", "Dammam"],
        ["EMP-1002", "Bilal Khan", "Chief Surveyor", "2456677890", "2023-10-01", "Pakistani", "9", "Survey", "Zuluf", "Jubail"],
        ["EMP-1003", "Rashid Ali", "Storekeeper", "2456677891", "2022-07-11", "Bangladeshi", "8", "Stores", "Warehouse", "Dammam"],
    ]


def import_employees(db: Database, records: Sequence[dict[str, Any]], source: str = "") -> tuple[int, int, int]:
    inserted = updated = skipped = 0
    for rec in records:
        row_id = None
        if str(rec.get("employee_id") or "").strip():
            found = db.one("SELECT id FROM employees WHERE employee_id=?", (str(rec.get("employee_id") or "").strip(),))
            if found is not None:
                row_id = int(found[0])
        if row_id is None and str(rec.get("iqama_id") or "").strip():
            found = db.one("SELECT id FROM employees WHERE iqama_id=?", (str(rec.get("iqama_id") or "").strip(),))
            if found is not None:
                row_id = int(found[0])
        try:
            if row_id is None:
                save_employee(db, dict(rec))
                inserted += 1
            else:
                save_employee(db, dict(rec), row_id)
                updated += 1
        except Exception:
            skipped += 1
    db.audit("IMPORTED", "employees", "", f"{inserted} inserted, {updated} updated, {skipped} skipped from {source}")
    return inserted, updated, skipped
